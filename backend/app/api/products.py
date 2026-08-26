from fastapi import APIRouter, Depends, HTTPException, status, UploadFile, File
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.future import select
from typing import List
from uuid import UUID
import csv
import openpyxl
import io

from app.database import get_db
from app.api.deps import get_current_user
from app.models.user import User
from app.models.business import Business
from app.models.product import Product
from app.schemas.product import ProductCreate, ProductRead, ProductUpdate

router = APIRouter()

async def get_user_business(business_id: UUID, current_user: User, db: AsyncSession) -> Business:
    result = await db.execute(
        select(Business).where(
            Business.id == business_id,
            Business.owner_id == current_user.id
        )
    )
    business = result.scalars().first()
    if not business:
        raise HTTPException(status_code=404, detail="Business not found")
    return business

@router.post("/{business_id}/products/upload/parse")
async def parse_upload(
    business_id: UUID,
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    await get_user_business(business_id, current_user, db)
    
    contents = await file.read()
    data = []
    headers = []
    
    try:
        if file.filename.endswith(".csv"):
            text = contents.decode("utf-8-sig")
            reader = csv.DictReader(io.StringIO(text))
            headers = reader.fieldnames or []
            data = [row for row in reader]
        elif file.filename.endswith(".xlsx"):
            wb = openpyxl.load_workbook(filename=io.BytesIO(contents), data_only=True)
            sheet = wb.active
            rows = list(sheet.iter_rows(values_only=True))
            if rows:
                headers = [str(h) if h is not None else "" for h in rows[0]]
                for row in rows[1:]:
                    data.append({headers[i]: (row[i] if row[i] is not None else "") for i in range(len(headers))})
        elif file.filename.endswith(".pdf"):
            import pdfplumber
            from app.services.llm_parser import parse_text_with_llm
            with pdfplumber.open(io.BytesIO(contents)) as pdf:
                raw_text = "\n".join([page.extract_text() for page in pdf.pages if page.extract_text()])
            products = parse_text_with_llm(raw_text)
            if not products:
                raise HTTPException(status_code=400, detail="Could not extract products via LLM from PDF.")
            headers = ["name", "price", "description", "quantity", "category"]
            return {"headers": headers, "data": products}
            
        elif file.filename.endswith(".docx"):
            import docx
            from app.services.llm_parser import parse_text_with_llm
            doc = docx.Document(io.BytesIO(contents))
            raw_text = "\n".join([p.text for p in doc.paragraphs])
            products = parse_text_with_llm(raw_text)
            if not products:
                raise HTTPException(status_code=400, detail="Could not extract products via LLM from DOCX.")
            headers = ["name", "price", "description", "quantity", "category"]
            return {"headers": headers, "data": products}
            
        else:
            raise HTTPException(status_code=400, detail="Unsupported file format. Please upload .csv, .xlsx, .pdf, or .docx")
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Error parsing file: {str(e)}")
        
    return {"headers": headers, "data": data}

@router.post("/{business_id}/products/bulk", status_code=status.HTTP_201_CREATED)
async def bulk_create_products(
    business_id: UUID,
    products_in: List[ProductCreate],
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    await get_user_business(business_id, current_user, db)
    
    new_products = []
    for p in products_in:
        prod = Product(
            **p.model_dump(),
            business_id=business_id,
        )
        prod.status = "draft"
        new_products.append(prod)
        
    db.add_all(new_products)
    await db.commit()
    
    return {"msg": f"Successfully imported {len(new_products)} products"}

@router.post("/{business_id}/products", response_model=ProductRead, status_code=status.HTTP_201_CREATED)
async def create_product(
    business_id: UUID,
    product_in: ProductCreate,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    await get_user_business(business_id, current_user, db)
    product = Product(
        **product_in.model_dump(),
        business_id=business_id
    )
    db.add(product)
    await db.commit()
    await db.refresh(product)
    return product

@router.get("/{business_id}/products", response_model=List[ProductRead])
async def read_products(
    business_id: UUID,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    await get_user_business(business_id, current_user, db)
    result = await db.execute(
        select(Product).where(Product.business_id == business_id)
    )
    return result.scalars().all()

@router.put("/{business_id}/products/{product_id}", response_model=ProductRead)
async def update_product(
    business_id: UUID,
    product_id: UUID,
    product_in: ProductUpdate,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    await get_user_business(business_id, current_user, db)
    
    result = await db.execute(
        select(Product).where(
            Product.id == product_id,
            Product.business_id == business_id
        )
    )
    product = result.scalars().first()
    if not product:
        raise HTTPException(status_code=404, detail="Product not found")

    update_data = product_in.model_dump(exclude_unset=True)
    for field, value in update_data.items():
        setattr(product, field, value)
        
    await db.commit()
    await db.refresh(product)
    return product

@router.delete("/{business_id}/products/{product_id}", status_code=status.HTTP_204_NO_CONTENT)
async def delete_product(
    business_id: UUID,
    product_id: UUID,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    await get_user_business(business_id, current_user, db)
    
    result = await db.execute(
        select(Product).where(
            Product.id == product_id,
            Product.business_id == business_id
        )
    )
    product = result.scalars().first()
    if not product:
        raise HTTPException(status_code=404, detail="Product not found")

    await db.delete(product)
    await db.commit()
    return None
